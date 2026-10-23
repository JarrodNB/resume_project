"""Build synthetic PDF/DOCX import samples from the labelled parsing cases.

Run ``python -m evals.resume_import_fixtures /tmp/resume-import-fixtures`` in the AI
container. Samples contain invented employers and accomplishments only. Wrapped-case
blank separators become actual page breaks; labels remain evaluation-only data.
"""

from pathlib import Path

import yaml


def build_pdf(pages):
    from io import BytesIO

    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        stream = DecodedStreamObject()
        encoded = []
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            encoded.append(b"(" + escaped.encode("cp1252") + b") Tj T*")
        stream.set_data(b"BT /F1 12 Tf 14 TL 50 700 Td " + b" ".join(encoded) + b" ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def build_docx(pages):
    from io import BytesIO

    from docx import Document

    document = Document()
    for index, lines in enumerate(pages):
        if index:
            document.add_page_break()
        for line in lines:
            if line.startswith("• "):
                document.add_paragraph(line[2:], style="List Bullet")
            else:
                document.add_paragraph(line)
    output = BytesIO()
    document.save(output)
    return output.getvalue()



FIXTURE_CASE_NAMES = (
    "chronological_backend_engineer", "grouped_shared_stack",
    "grouped_mixed_role_and_shared_lists", "separate_same_company_blocks",
    "grouped_wrapped_acronyms_and_blank_pages",
)


def fixture_pages(case):
    """Turn the wrapped sample's intentional blank separators into physical pages."""
    text = case["inputs"]
    chunks = text.split("\n\n") if case["name"] == "grouped_wrapped_acronyms_and_blank_pages" else [text]
    return [chunk.splitlines() for chunk in chunks if chunk.strip()]


def build_fixtures(destination: Path):
    """Write human-uploadable documents and their source/expected labels together."""
    source = Path(__file__).parent / "datasets" / "resume_parsing_cases.yaml"
    cases = yaml.safe_load(source.read_text())["cases"]
    destination.mkdir(parents=True, exist_ok=True)
    for case in cases:
        if case["name"] not in FIXTURE_CASE_NAMES:
            continue
        pages = fixture_pages(case)
        for extension, builder in [("pdf", build_pdf), ("docx", build_docx)]:
            (destination / f"{case['name']}.{extension}").write_bytes(builder(pages))
        (destination / f"{case['name']}.txt").write_text(case["inputs"])
        (destination / f"{case['name']}.expected.yaml").write_text(yaml.safe_dump(case["expected_output"], sort_keys=False))


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("destination", type=Path)
    build_fixtures(parser.parse_args().destination)
