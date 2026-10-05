"""
Description: Test doubles and builders shared by the pytest suite.
Index:
    - FakeFileStorage: serves file bytes from a dict instead of S3
    - build_docx_bytes: builds a small DOCX file in memory
    - build_gap_analysis_context: a small job and career library for gap analysis tests
    - FakeGapAnalysisContextReader: serves gap analysis contexts from a dict instead of Postgres
    - FakeJobSite: answers HTTP requests from a dict of URL -> HTML (404 for anything else)
"""

from datetime import date
from io import BytesIO
from uuid import UUID

import httpx
from docx import Document

from schemas.gap_analysis_context import (
    CareerLibrary,
    GapAnalysisContext,
    JobToAnalyze,
    LibraryEducation,
    LibraryExperience,
    LibraryHighlight,
    LibraryProject,
    LibrarySkill,
)


class FakeFileStorage:
    """Stands in for S3FileStorage."""

    def __init__(self) -> None:
        """Starts empty; tests add files with `files[key] = bytes`."""
        self.files: dict[str, bytes] = {}

    def download_file(self, key: str) -> bytes:
        """Returns the stored bytes, raising KeyError for unknown keys like S3 would."""
        return self.files[key]


def build_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    """Builds a DOCX file with the given paragraphs and an optional table."""
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row, values in zip(table.rows, table_rows, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_gap_analysis_context(
    requirements: list[str] | None = None, nice_to_haves: list[str] | None = None
) -> GapAnalysisContext:
    """A small job and career library: one ongoing role with two highlights, a project, a
    degree, and two skills (prompt labels E1, H1, H2, P1, ED1, S1, S2). The profile summary
    is background text, not a label."""
    return GapAnalysisContext(
        job=JobToAnalyze(
            id="job-1",
            title="Platform Engineer",
            company="Acme",
            requirements=["5+ years of Python", "PostgreSQL"]
            if requirements is None
            else requirements,
            nice_to_haves=["Kubernetes"] if nice_to_haves is None else nice_to_haves,
        ),
        career_library=CareerLibrary(
            summary="Backend engineer.",
            experiences=[
                LibraryExperience(
                    id="experience-1",
                    company="Now Co",
                    title="Senior Engineer",
                    start_date=date(2018, 1, 1),
                    highlights=[
                        LibraryHighlight(id="highlight-1", content="Built Python APIs"),
                        LibraryHighlight(id="highlight-2", content="Tuned Postgres queries"),
                    ],
                )
            ],
            projects=[LibraryProject(id="project-1", name="Deploy bot", role="Author")],
            educations=[
                LibraryEducation(
                    id="education-1",
                    institution="State University",
                    degree="BSc",
                    field_of_study="Computer Science",
                )
            ],
            skills=[
                LibrarySkill(id="skill-1", name="Python", category="Languages"),
                LibrarySkill(id="skill-2", name="PostgreSQL"),
            ],
        ),
    )


class FakeGapAnalysisContextReader:
    """Stands in for the database reader: serves contexts keyed by (user_id, job_id)."""

    def __init__(self) -> None:
        """Starts empty; tests add contexts with `contexts[(user_id, job_id)] = context`."""
        self.contexts: dict[tuple[UUID, UUID], GapAnalysisContext] = {}

    def __call__(self, user_id: UUID, job_id: UUID) -> GapAnalysisContext | None:
        """Returns the stored context, or None like a job the user doesn't have."""
        return self.contexts.get((user_id, job_id))


class FakeJobSite:
    """Stands in for job sites, through httpx.MockTransport."""

    def __init__(self) -> None:
        """Starts empty; tests add pages with `pages[url] = html` or `statuses[url] = code`."""
        self.pages: dict[str, str] = {}
        self.statuses: dict[str, int] = {}
        self.requested_urls: list[str] = []

    def handle_request(self, request: httpx.Request) -> httpx.Response:
        """Returns the configured HTML or status for the URL, else 404."""
        url = str(request.url)
        self.requested_urls.append(url)
        if url in self.statuses:
            return httpx.Response(self.statuses[url])
        if url in self.pages:
            return httpx.Response(200, html=self.pages[url])
        return httpx.Response(404)

    def build_client(self) -> httpx.Client:
        """An httpx client whose requests this fake answers."""
        return httpx.Client(transport=httpx.MockTransport(self.handle_request))
