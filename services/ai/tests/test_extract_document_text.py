"""Check resume text extraction, document order, links, lists, and import limits."""

from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from documents.extract_document_text import (
    DOCX_CONTENT_TYPE,
    DocumentTooLongError,
    UnsupportedDocumentError,
    _extract_pdf_page,
    extract_document_text,
)
from tests.fakes import build_docx_bytes
from text_limits import MAX_PDF_PAGES, MAX_RESUME_TEXT


def test_docx_text_includes_paragraphs_and_table_cells() -> None:
    """Resumes laid out with tables still yield all their text."""
    file_bytes = build_docx_bytes(["Jane Doe", "Experience"], table_rows=[["Acme", "2021 - 2023"]])

    text = extract_document_text(file_bytes, DOCX_CONTENT_TYPE)

    assert text.splitlines() == ["Jane Doe", "Experience", "Acme", "2021 - 2023"]


def test_docx_text_includes_nested_tables_in_cell_order() -> None:
    """Nested resume layouts keep their details between surrounding cell paragraphs."""
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Experience"
    nested_cell = cell.add_table(rows=1, cols=1).cell(0, 0)
    nested_cell.paragraphs[0].text = "Example Company"
    nested_cell.add_table(rows=1, cols=1).cell(0, 0).text = "Built Python APIs"
    cell.add_paragraph("Education")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert [line for line in text.splitlines() if line] == [
        "Experience", "Example Company", "Built Python APIs", "Education",
    ]


@pytest.mark.parametrize("vertical", [False, True])
def test_docx_merged_cells_are_read_once(vertical: bool) -> None:
    """Merged resume headings and descriptions must not become duplicate evidence."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "Software Engineer"
    table.cell(1, 0).text = "Built Python APIs"
    table.cell(1, 1).text = "Education"
    table.cell(0, 0).merge(table.cell(1, 0) if vertical else table.cell(0, 1))
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    expected = (
        ["Experience", "Built Python APIs", "Software Engineer", "Education"]
        if vertical else ["Experience", "Software Engineer", "Built Python APIs", "Education"]
    )
    assert text.splitlines() == expected


def test_docx_text_preserves_list_bullet_markers_for_exact_highlight_mapping() -> None:
    document = Document()
    document.add_paragraph("Built a reliable API", style="List Bullet")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert text == "• Built a reliable API"


def test_docx_text_keeps_tables_between_surrounding_paragraphs() -> None:
    """Table-based job details stay under their original section heading."""
    document = Document()
    document.add_paragraph("Experience")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Example Company"
    table.cell(0, 1).text = "Software Engineer"
    document.add_paragraph("Education")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Example University"
    document.add_paragraph("Skills")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert text.splitlines() == [
        "Experience",
        "Example Company",
        "Software Engineer",
        "Education",
        "Example University",
        "Skills",
    ]


def test_docx_text_detects_numbering_metadata_on_normal_style_paragraph() -> None:
    document = Document()
    paragraph = document.add_paragraph("Built a reliable API")
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number_id = OxmlElement("w:numId")
    number_id.set(qn("w:val"), "1")
    num_properties.append(level)
    num_properties.append(number_id)
    paragraph._p.get_or_add_pPr().append(num_properties)
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert text == "• Built a reliable API"


@pytest.mark.parametrize("style", ["Normal", "List Bullet", "List Number"])
def test_docx_explicitly_cancelled_numbering_has_no_list_marker(style: str) -> None:
    """OOXML numId zero removes direct or inherited list formatting."""
    document = Document()
    paragraph = document.add_paragraph("Software Engineer", style=style)
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    buffer = BytesIO()
    document.save(buffer)

    assert extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE) == "Software Engineer"


def test_docx_style_can_cancel_inherited_numbering() -> None:
    """A derived style can turn off the list formatting of its base style."""
    from docx.enum.style import WD_STYLE_TYPE

    document = Document()
    style = document.styles.add_style("Plain role heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["List Bullet"]
    style._element.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
    document.add_paragraph("Software Engineer", style=style)
    buffer = BytesIO()
    document.save(buffer)

    assert extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE) == "Software Engineer"


@pytest.mark.parametrize("style", ["List Bullet", "List Number"])
def test_docx_partial_numbering_properties_still_inherit_list(style: str) -> None:
    """An indentation-level override alone does not remove a style's list ID."""
    document = Document()
    paragraph = document.add_paragraph("Built APIs", style=style)
    paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl().val = 0
    buffer = BytesIO()
    document.save(buffer)

    marker = "1." if style == "List Number" else "•"
    assert extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE) == f"{marker} Built APIs"


def test_docx_text_includes_hyperlink_label_and_url() -> None:
    """A labelled DOCX hyperlink keeps both the visible text and destination."""
    document = Document()
    paragraph = document.add_paragraph("Find me on ")
    relationship_id = paragraph.part.relate_to(
        "https://github.com/jane",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    text_node = OxmlElement("w:t")
    text_node.text = "GitHub"
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert "Find me on GitHub" in text
    assert '[hyperlink label="GitHub" url="https://github.com/jane"]' in text

def test_pdf_text_uses_only_the_link_label_from_a_broad_contact_fragment() -> None:
    """A URI annotation does not use the entire contact line as its visible label."""
    contact_line = "Atlanta, Georgia · 770 - 202 - 0267 JarrodNBailey@gmail.com · LinkedIn ·"

    class Reference:
        """A small indirect-object stand-in for one PDF annotation."""

        def get_object(self) -> dict[str, object]:
            """Returns an external LinkedIn annotation over part of the contact line."""
            return {
                "/A": {"/URI": "https://www.linkedin.com/in/jarrod-bailey-68w"},
                "/Rect": [300, 8, 360, 24],
            }

    class Page(dict[str, object]):
        """A page whose extractor emits the whole contact line as one text fragment."""

        def __init__(self) -> None:
            super().__init__({"/Annots": [Reference()]})

        def extract_text(self, visitor_text: object) -> str:
            """Feeds the visitor the broad fragment that previously became the label."""
            visitor_text(contact_line, [1, 0, 0, 1, 0, 0], [1, 0, 0, 1, 10, 10], None, 10)
            return contact_line

    text = _extract_pdf_page(Page())

    assert (
        '[hyperlink label="LinkedIn" '
        'url="https://www.linkedin.com/in/jarrod-bailey-68w"]'
    ) in text
    assert 'hyperlink label="Atlanta, Georgia' not in text


def test_empty_docx_raises() -> None:
    """A document with no text can't be parsed."""
    with pytest.raises(UnsupportedDocumentError):
        extract_document_text(build_docx_bytes([]), DOCX_CONTENT_TYPE)


def test_docx_over_the_text_cap_is_too_long() -> None:
    """A real upload over the import cap is a length error."""
    with pytest.raises(DocumentTooLongError):
        extract_document_text(build_docx_bytes(["a" * (MAX_RESUME_TEXT + 1)]), DOCX_CONTENT_TYPE)


def test_docx_at_the_text_cap_is_accepted() -> None:
    """The maximum allowed text length remains importable."""
    text = "a" * MAX_RESUME_TEXT
    assert extract_document_text(build_docx_bytes([text]), DOCX_CONTENT_TYPE) == text


def test_pdf_over_the_page_cap_is_too_long() -> None:
    """Real PDF pages over the cap are rejected before empty-text validation."""
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(MAX_PDF_PAGES + 1):
        writer.add_blank_page(width=300, height=300)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(DocumentTooLongError):
        extract_document_text(buffer.getvalue(), "application/pdf")


@pytest.mark.parametrize("extra_characters", [0, 1])
def test_pdf_text_cap_applies_across_pages(extra_characters: int) -> None:
    """The import cap includes all pages and separators, with the exact limit allowed."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    page_texts = [
        "a" * (MAX_RESUME_TEXT // 2 - 1),
        "b" * (MAX_RESUME_TEXT // 2 - 1 + extra_characters),
    ]
    writer = PdfWriter()
    for text in page_texts:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 20 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)

    if extra_characters:
        with pytest.raises(DocumentTooLongError):
            extract_document_text(buffer.getvalue(), "application/pdf")
    else:
        assert extract_document_text(buffer.getvalue(), "application/pdf") == "\n\n".join(page_texts)


def test_unsupported_content_type_raises() -> None:
    """Only PDF and DOCX are supported."""
    with pytest.raises(UnsupportedDocumentError):
        extract_document_text(b"hello", "text/plain")


@pytest.mark.parametrize("uri", ["https://[broken/profile", "https://example.com\uff0f.evil/profile"])
def test_pdf_malformed_link_does_not_abort_extraction(uri: str) -> None:
    """An invalid destination still leaves its annotation available as source text."""
    from pypdf import PdfWriter
    from pypdf.annotations import Link

    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.add_annotation(0, Link(rect=(10, 10, 100, 30), url=uri))
    buffer = BytesIO()
    writer.write(buffer)

    text = extract_document_text(buffer.getvalue(), "application/pdf")

    assert f'url="{uri}"' in text


def test_pdf_indirect_annotation_array_preserves_links() -> None:
    """The PDF annotation array may live in a separate indirect object."""
    from pypdf import PdfWriter
    from pypdf.annotations import Link
    from pypdf.generic import NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    writer.add_annotation(0, Link(rect=(10, 10, 100, 30), url="https://github.com/example"))
    page[NameObject("/Annots")] = writer._add_object(page["/Annots"])
    buffer = BytesIO()
    writer.write(buffer)

    text = extract_document_text(buffer.getvalue(), "application/pdf")

    assert '[hyperlink label="GitHub" url="https://github.com/example"]' in text


@pytest.mark.parametrize("indirect", [False, True])
def test_pdf_null_annotations_are_treated_as_absent(indirect: bool) -> None:
    """A null optional entry is equivalent to an absent entry in a PDF dictionary."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NullObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 20 Td (Software Engineer) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    null = NullObject()
    page[NameObject("/Annots")] = writer._add_object(null) if indirect else null
    buffer = BytesIO()
    writer.write(buffer)

    assert extract_document_text(buffer.getvalue(), "application/pdf") == "Software Engineer"


def test_malformed_docx_xml_is_an_unreadable_document() -> None:
    """A corrupt XML part uses the import handler's ordinary unreadable-file path."""
    from zipfile import ZipFile

    original = build_docx_bytes(["Software Engineer"])
    buffer = BytesIO()
    with ZipFile(BytesIO(original)) as source, ZipFile(buffer, "w") as destination:
        for entry in source.infolist():
            destination.writestr(
                entry, b"<broken" if entry.filename == "word/document.xml" else source.read(entry)
            )

    with pytest.raises(UnsupportedDocumentError, match="could not be read"):
        extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)


@pytest.mark.parametrize("explicit_style", [None, "MissingStyle"])
def test_docx_without_a_default_style_keeps_plain_text(explicit_style: str | None) -> None:
    """Optional or unresolved style metadata must not prevent reading a paragraph."""
    document = Document()
    paragraph = document.add_paragraph("Software Engineer")
    document.styles["Normal"].delete()
    if explicit_style is not None:
        paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = explicit_style
    buffer = BytesIO()
    document.save(buffer)

    assert extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE) == "Software Engineer"


def test_docx_keeps_header_and_footer_resume_content() -> None:
    """Contact details and portfolio labels in margins remain available to the parser."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe jane@example.com"
    document.add_paragraph("Software Engineer")
    document.sections[0].footer.paragraphs[0].text = "Portfolio: https://example.com"
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert text.splitlines() == [
        "Jane Doe jane@example.com", "Software Engineer", "Portfolio: https://example.com",
    ]


def test_docx_inherited_headers_and_footers_are_read_once() -> None:
    """Shared margin parts must not duplicate evidence across document sections."""
    from docx.enum.section import WD_SECTION

    document = Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe"
    document.sections[0].footer.paragraphs[0].text = "Portfolio"
    document.add_paragraph("Engineer")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Education")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert [line for line in text.splitlines() if line] == [
        "Jane Doe", "Engineer", "Education", "Portfolio",
    ]


@pytest.mark.parametrize("variant", ["first_page", "even_page"])
@pytest.mark.parametrize("enabled", [False, True])
def test_docx_margin_variants_follow_document_settings(variant: str, enabled: bool) -> None:
    """Unused margin variants are not source evidence merely because their parts exist."""
    document = Document()
    section = document.sections[0]
    section.different_first_page_header_footer = enabled
    document.settings.odd_and_even_pages_header_footer = enabled
    getattr(section, f"{variant}_header").paragraphs[0].text = "Variant contact"
    getattr(section, f"{variant}_footer").paragraphs[0].text = "Variant portfolio"
    document.add_paragraph("Engineer")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert ("Variant contact" in text) == enabled
    assert ("Variant portfolio" in text) == enabled
    assert "Engineer" in text


def test_docx_header_tables_preserve_hyperlink_labels_and_targets() -> None:
    """Contact layouts in header tables use the same text/link extraction as body tables."""
    from docx.shared import Inches

    document = Document()
    paragraph = document.sections[0].header.add_table(1, 1, Inches(6)).cell(0, 0).paragraphs[0]
    relationship_id = paragraph.part.relate_to(
        "https://github.com/example",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = "GitHub"
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.add_paragraph("Engineer")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert [line for line in text.splitlines() if line] == [
        "GitHub", '[hyperlink label="GitHub" url="https://github.com/example"]', "Engineer",
    ]


@pytest.mark.parametrize("extra_characters", [0, 1])
def test_docx_text_cap_includes_margin_content(extra_characters: int) -> None:
    """The same import cap applies to combined margin and body text."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "h"
    document.add_paragraph("b" * (MAX_RESUME_TEXT - 2 + extra_characters))
    buffer = BytesIO()
    document.save(buffer)

    if extra_characters:
        with pytest.raises(DocumentTooLongError):
            extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)
    else:
        assert len(extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)) == MAX_RESUME_TEXT
