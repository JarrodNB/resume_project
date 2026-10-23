"""Keep synthetic resume bullets intact while respecting employment and section boundaries.

These tests use the public source selector and real file extractors. Model responses are
synthetic selections; they test reconstruction rather than model interpretation quality.
"""

import re

import pytest

from documents.experience_source import ExperienceSource
from evals.resume_import_fixtures import build_docx, build_pdf


def source_highlights(text):
    source = ExperienceSource(text)
    ids = re.findall(r"\[(B\d+)\]", source.annotated_text)
    output = source.output_type.model_validate({
        "experiences": [{"company": "Example", "title": "Engineer", "highlight_ids": ids}],
    })
    return source.reconstruct(output)["experiences"][0]["roles"][0]["highlights"]


def test_stacked_role_header_with_dates_does_not_enter_previous_bullet():
    text = (
        "WORK EXPERIENCE\nExample | Remote\nEngineer | January 2025 – September 2026\n"
        "• Built a payment system used across three regions.\n"
        "Software Developer | 2019 – 2024\n"
        "• Reduced deployment failures.\n"
    )
    assert source_highlights(text) == [
        "Built a payment system used across three regions.", "Reduced deployment failures.",
    ]


def test_acronym_and_blank_line_continuations_stay_in_one_bullet():
    text = (
        "WORK EXPERIENCE\nEngineer | 2020 – Present\n"
        "• Integrated the platform with\nAWS SQS\n\n"
        "to deliver reliable notifications.\n"
        "• Built the reporting service.\nSKILLS\n• Python\n"
    )
    assert source_highlights(text) == [
        "Integrated the platform with\nAWS SQS\nto deliver reliable notifications.",
        "Built the reporting service.",
    ]


def test_blank_separators_do_not_hide_real_boundaries():
    text = (
        "WORK EXPERIENCE\n• Built the first service.\n\n"
        "Software Engineer\n\nExample\n\n2020 – Present\n"
        "• Built the next service.\n\nEDUCATION\n• Studied computing.\n"
    )
    assert source_highlights(text) == ["Built the first service.", "Built the next service."]


@pytest.mark.parametrize("file_format", ["pdf", "docx"])
def test_uploaded_files_preserve_stacked_bullets_through_final_parser(
    file_format, offline_model_resolution,
):
    from pydantic_ai.messages import ModelResponse, ToolCallPart
    from pydantic_ai.models.function import FunctionModel

    from agents.resume_parsing_agent import parse_resume_text
    from documents.extract_document_text import (
        DOCX_CONTENT_TYPE,
        PDF_CONTENT_TYPE,
        extract_document_text,
    )

    pages = [
        [
            "WORK EXPERIENCE", "Earlier Company", "Engineer | 2017 – 2018",
            "• Built a payment system used across three regions.",
            "Example | Remote", "AI Integrations Engineer | 2025 – 2026",
            "Software Developer | 2019 – 2024",
            "• Integrated the platform with", "AWS SQS",
        ],
        [
            "to deliver reliable notifications.", "",
            "• Built the reporting service.", "EDUCATION", "• Studied computing.",
        ],
    ]
    if file_format == "pdf":
        file_bytes, content_type = build_pdf(pages), PDF_CONTENT_TYPE
    else:
        file_bytes, content_type = build_docx(pages), DOCX_CONTENT_TYPE
    source = extract_document_text(file_bytes, content_type)
    expected = [
        "Built a payment system used across three regions.",
        "Integrated the platform with\nAWS SQS\nto deliver reliable notifications.",
        "Built the reporting service.",
    ]
    assert source_highlights(source) == expected

    async def answer(messages, info):
        if "experiences" in info.output_tools[0].parameters_json_schema["properties"]:
            args = {"experiences": [
                {"company": "Earlier Company", "title": "Engineer", "highlight_ids": ["B1"]},
                {"company": "Example", "title": "Software Developer",
                 "highlight_ids": ["B3", "B2"]},
            ]}
        else:
            args = {}
        return ModelResponse(parts=[ToolCallPart(info.output_tools[0].name, args)])

    offline_model_resolution("openai:gpt-5-mini", FunctionModel(answer))
    result = parse_resume_text(source, "openai:gpt-5-mini", "low", 5, "openai:gpt-5-mini")
    assert [role.roles[0].highlights for role in result.parsed_resume.experiences] == [
        expected[:1], expected[1:],
    ]


def test_short_complete_sentence_keeps_its_following_paragraph():
    text = (
        "WORK EXPERIENCE\nEngineer | 2020 – Present\n"
        "• Built APIs.\n\nThis reduced outages.\n"
        "SKILLS\n• Python\n"
    )
    assert source_highlights(text) == ["Built APIs.\nThis reduced outages."]


@pytest.mark.parametrize("separator", ["", "\n"])
@pytest.mark.parametrize("punctuation", ["", "."])
@pytest.mark.parametrize("heading", ["PROFESSIONAL DEVELOPMENT", "Professional Development"])
def test_professional_development_stays_outside_highlights(separator, punctuation, heading):
    text = (
        "WORK EXPERIENCE\nEngineer | 2020 – Present\n"
        f"• Built a reliable payment service{punctuation}\n{separator}"
        f"{heading}\nCompleted training in distributed systems.\n"
    )
    assert source_highlights(text) == [f"Built a reliable payment service{punctuation}"]


@pytest.mark.parametrize("experience_heading", ["WORK EXPERIENCE\n", ""])
def test_professional_development_bullets_and_dated_courses_are_not_employment(experience_heading):
    text = (
        f"{experience_heading}Engineer | 2020 – Present\n"
        "• Built payment services\n\nPROFESSIONAL DEVELOPMENT\n"
        "• Completed distributed systems training\n"
        "Systems Course | 2023 – 2024\n• Studied message delivery\n"
    )
    assert source_highlights(text) == ["Built payment services"]


@pytest.mark.parametrize("separator", ["", "\n"])
def test_acronym_continuation_after_a_complete_sentence_stays_in_its_bullet(separator):
    text = (
        "WORK EXPERIENCE\nEngineer | 2020 – Present\n"
        f"• Built reliable services.\n{separator}AWS SQS\n"
        "handled our notification queue.\n"
    )
    assert source_highlights(text) == [
        "Built reliable services.\nAWS SQS\nhandled our notification queue.",
    ]


def test_unlabelled_list_after_stacked_dated_titles_cannot_belong_to_one_role():
    source = ExperienceSource(
        'EXPERIENCE\nExample | Remote\n'
        'Senior Engineer | 2025 – 2026\nEngineer | 2019 – 2024\n'
        'Built products across both positions.\n• Shipped the platform.\n'
    )
    with pytest.raises(ValueError, match='shared_highlight_ids'):
        source.output_type.model_validate({'experiences': [{
            'company': 'Example', 'roles': [
                {'title': 'Senior Engineer', 'highlight_ids': ['B1']},
                {'title': 'Engineer', 'highlight_ids': []},
            ],
        }]})


def test_docx_topic_headings_do_not_become_highlight_continuations():
    from io import BytesIO

    from docx import Document

    from documents.extract_document_text import DOCX_CONTENT_TYPE, extract_document_text
    document = Document()
    document.add_heading('EXPERIENCE', 1)
    document.add_paragraph('Example | Remote')
    document.add_paragraph('Engineer | 2020 – Present')
    document.add_paragraph('Shipped the platform.', style='List Bullet')
    document.add_heading('Evaluation & Guardrails', 4)
    document.add_paragraph('Reduced errors.', style='List Bullet')
    buffer = BytesIO()
    document.save(buffer)
    assert source_highlights(extract_document_text(buffer.getvalue(), DOCX_CONTENT_TYPE)) == [
        'Shipped the platform.', 'Reduced errors.',
    ]


def test_stacked_titles_keep_valid_shared_bullets():
    source = ExperienceSource('EXPERIENCE\nExample | Remote\nSenior Engineer | 2025 – 2026\nEngineer | 2019 – 2024\n• Shipped the platform.\n')
    output = source.output_type.model_validate({'experiences': [{'company': 'Example', 'roles': [
        {'title': 'Senior Engineer'}, {'title': 'Engineer'},
    ], 'shared_highlight_ids': ['B1']}]})
    assert source.reconstruct(output)['experiences'][0]['shared_highlights'] == ['Shipped the platform.']


def test_bullets_between_dated_roles_remain_scoped():
    source = ExperienceSource(
        'EXPERIENCE\nExample | Remote\nSenior Engineer | 2025 – 2026\n'
        '• Shipped the platform.\nEngineer | 2019 – 2024\n• Built the service.\n'
    )
    output = source.output_type.model_validate({'experiences': [{'company': 'Example', 'roles': [
        {'title': 'Senior Engineer', 'highlight_ids': ['B1']},
        {'title': 'Engineer', 'highlight_ids': ['B2']},
    ]}]})
    assert [role['highlights'] for role in source.reconstruct(output)['experiences'][0]['roles']] == [
        ['Shipped the platform.'], ['Built the service.'],
    ]
